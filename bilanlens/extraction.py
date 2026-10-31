from pathlib import Path
from datetime import datetime
import io
import re
import shutil

import pandas as pd

FORMATS_PRIS_EN_CHARGE = {".pdf", ".xlsx", ".xls", ".csv", ".txt", ".docx"}

MOTIF_DATE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b")

# Seuil de caractères par page en dessous duquel on considère la page
# comme un scan (image). On déclenche alors l'OCR si disponible.
OCR_MIN_CARACTERES_PAR_PAGE = 40


_CHEMINS_TESSERACT_REPLI = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    "/opt/homebrew/bin/tesseract",
    "/usr/local/bin/tesseract",
    "/usr/bin/tesseract",
]


def _localiser_binaire_tesseract():
    """Localise tesseract via PATH, puis via les chemins standards."""
    trouve = shutil.which("tesseract")
    if trouve:
        return trouve
    for candidat in _CHEMINS_TESSERACT_REPLI:
        if Path(candidat).exists():
            return candidat
    return None


def ocr_disponible() -> bool:
    """Retourne True si Tesseract et pytesseract sont disponibles."""
    if _localiser_binaire_tesseract() is None:
        return False
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return True


def _configurer_pytesseract(module):
    """Pointe pytesseract vers le binaire Tesseract trouvé."""
    binaire = _localiser_binaire_tesseract()
    if not binaire:
        return
    # pytesseract expose tesseract_cmd soit sur le module lui-même,
    # soit sur un attribut interne `pytesseract` selon la version.
    interne = getattr(module, "pytesseract", module)
    interne.tesseract_cmd = binaire


_CORRECTIONS_NOMBRES_OCR = str.maketrans({
    # Fautes OCR fréquentes dans les nombres. On ne corrige que les
    # chaînes qui ressemblent déjà à un nombre (voir _corriger_token_nombre_ocr).
    "O": "0",
    "o": "0",
    "l": "1",
    "I": "1",
    "S": "5",
    "B": "8",
})


_MOTIF_TOKEN_NOMBRE = re.compile(
    r"(?<![A-Za-zÀ-ÿ])"
    r"(?:\(\s*)?"
    r"[-+]?[0-9OolISB][0-9OolISB.,\s  ]*"
    r"(?:\s*\))?"
    r"(?![A-Za-zÀ-ÿ])"
)


def _grouper_mots_par_ligne(page, tolerance_y: int = 3) -> list[list[dict]]:
    """Regroupe les mots d'une page par ligne visuelle (tolérance verticale),
    trie chaque ligne par `x0`. Retourne `[[mots_ligne_0], [mots_ligne_1], ...]`.

    Helper partagé par `_pdf_page_text` (reconstruction du texte) et
    `_extraire_lignes_periodes_depuis_page_pdf` (extraction par coordonnées
    des montants périodiques). Une seule logique de bucketing.

    `tolerance_y` : 2-3 px pour un PDF natif ; 15-30 px pour une page OCR
    à 300 dpi (textbaselines décalés).
    """
    try:
        mots = page.extract_words(
            x_tolerance=2,
            y_tolerance=tolerance_y,
            keep_blank_chars=False,
        )
    except TypeError:
        # Shims de test ou pages OCR qui ignorent les kwargs.
        mots = page.extract_words()

    if not mots:
        return []

    # Fenêtre glissante sur les 3 dernières lignes vues, pour tolérer un
    # léger ordonnancement non monotone des mots en sortie de pdfplumber.
    lignes: list[dict] = []
    for mot in sorted(mots, key=lambda w: (w["top"], w["x0"])):
        cible = None
        for ligne in reversed(lignes[-3:]):
            if abs(ligne["top"] - mot["top"]) <= tolerance_y:
                cible = ligne
                break
        if cible is None:
            lignes.append({"top": mot["top"], "mots": [mot]})
        else:
            cible["mots"].append(mot)
    lignes.sort(key=lambda l: l["top"])
    return [sorted(l["mots"], key=lambda w: w["x0"]) for l in lignes]


def _pdf_page_text(
    page, tolerance_y: int = 3, facteur_colonne: float = 1.0
) -> str:
    """Reconstruit le texte d'une page pdfplumber à partir de la POSITION
    des mots, en préservant la séparation des colonnes.

    `page.extract_text()` colle parfois des colonnes voisines si leurs
    glyphes sont proches en x : une ligne « Immobilisations 85 230 110
    82 450 220 » peut devenir « Immobilisations 85230110 82450220 ». Ici,
    on insère un espace simple pour un petit écart (= espace de milliers
    « 3 500 ») et un double espace pour un grand écart (= nouvelle colonne).

    Seuil adaptatif : nouvelle colonne si l'écart dépasse
    `facteur_colonne * hauteur_mot` (hauteur = `bottom - top`, ≈ taille de
    police). Un seuil fixe en pixels tombait à faux dès qu'on changeait de
    taille de police ; la hauteur du glyphe est quasi-proportionnelle à la
    taille, donc le seuil suit automatiquement.
    """
    lignes_mots = _grouper_mots_par_ligne(page, tolerance_y=tolerance_y)
    if not lignes_mots:
        return ""

    textes: list[str] = []
    for mots_tries in lignes_mots:
        parties: list[str] = [mots_tries[0]["text"]]
        for prec, mot in zip(mots_tries, mots_tries[1:]):
            ecart = mot["x0"] - prec["x1"]
            # Hauteur du mot courant si disponible, sinon valeur par défaut
            # raisonnable (shims de test / mots OCR anciens sans `bottom`).
            hauteur = max(1.0, mot.get("bottom", mot["top"] + 10) - mot["top"])
            seuil = facteur_colonne * hauteur
            parties.append("  " if ecart > seuil else " ")
            parties.append(mot["text"])
        textes.append("".join(parties))
    return "\n".join(textes)


def _corriger_token_nombre_ocr(token: str) -> str:
    """Corrige O/l/I/S/B dans un token qui ressemble à un nombre OCR.

    N'altère rien si le token ne contient aucun chiffre : on évite de
    transformer un vrai mot (« Oui », « Lot »...) en nombre.
    """
    chiffres_presents = any(ch.isdigit() for ch in token)
    lettres_presentes = any(ch.isalpha() for ch in token)
    if not chiffres_presents or not lettres_presentes:
        return token
    return token.translate(_CORRECTIONS_NOMBRES_OCR)


def nettoyer_texte_ocr(texte: str) -> str:
    """Corrige les erreurs OCR fréquentes dans les nombres.

    Ne modifie pas les libellés ni n'invente de valeurs : seuls les
    tokens numériques comportant une lettre confusable (O, l, I, S, B)
    sont corrigés.
    """
    if not texte:
        return texte

    def _remplacer(correspondance):
        return _corriger_token_nombre_ocr(correspondance.group(0))

    nettoye = _MOTIF_TOKEN_NOMBRE.sub(_remplacer, texte)
    # Supprimer les espaces insécables parasites qui gênent le découpage
    # en colonnes (on garde les espaces normaux).
    nettoye = nettoye.replace(" ", " ").replace(" ", " ")
    return nettoye


def _ocr_pages(
    pdf_source,
    indices_cibles: list[int] | None = None,
    langues: str = "fra+eng",
    avec_mots: bool = False,
):
    """OCR via PyMuPDF + Tesseract, limité aux pages demandées.

    `pdf_source` : `Path | str | bytes` (octets bruts d'un PDF en mémoire).
    `indices_cibles` : numéros de pages 1-based à OCR. `None` = toutes les
    pages (lent). Vide = retourne `{}` directement.
    `avec_mots` : si `True`, retourne `{page_num: (texte, mots)}` avec les
    coordonnées (`text/x0/x1/top/bottom`) de chaque mot ; si `False`
    (défaut, usage chunking), retourne `{page_num: texte}`.

    Retourne `{}` quand Tesseract/PyMuPDF/Pillow sont absents ou si l'OCR
    échoue. Boucle unique partagée avec `_ocr_pdf_octets`.
    """
    if not ocr_disponible():
        return {}
    if indices_cibles is not None and not indices_cibles:
        return {}

    try:
        import pymupdf  # type: ignore
    except ImportError:
        try:
            # `fitz` est l'ancien nom du module importé après `pip install pymupdf`.
            # Les versions récentes exposent les deux noms ; les anciennes
            # (< 1.24) n'exposaient que `fitz`.
            import fitz as pymupdf  # type: ignore
        except ImportError:
            return {}

    try:
        import pytesseract
        from pytesseract import Output
        from PIL import Image
    except ImportError:
        return {}

    _configurer_pytesseract(pytesseract)

    cibles = set(indices_cibles) if indices_cibles is not None else None
    resultats: dict = {}

    if isinstance(pdf_source, (bytes, bytearray)):
        doc_ctx = pymupdf.open(stream=bytes(pdf_source), filetype="pdf")
    else:
        doc_ctx = pymupdf.open(pdf_source)

    try:
        with doc_ctx as doc:
            matrice = pymupdf.Matrix(300 / 72, 300 / 72)
            for i, page in enumerate(doc, start=1):
                if cibles is not None and i not in cibles:
                    continue
                pix = page.get_pixmap(matrix=matrice, alpha=False)
                img = Image.frombytes(
                    "RGB", (pix.width, pix.height), pix.samples
                )
                try:
                    texte = pytesseract.image_to_string(img, lang=langues)
                    donnees = (
                        pytesseract.image_to_data(
                            img, lang=langues, output_type=Output.DICT
                        )
                        if avec_mots
                        else None
                    )
                except pytesseract.TesseractError:
                    # Pack `fra` absent sur certains CI : repli sur `eng`.
                    texte = pytesseract.image_to_string(img, lang="eng")
                    donnees = (
                        pytesseract.image_to_data(
                            img, lang="eng", output_type=Output.DICT
                        )
                        if avec_mots
                        else None
                    )
                if not avec_mots:
                    resultats[i] = texte
                    continue
                # Mots au format pdfplumber (text/x0/x1/top/bottom), pour
                # pouvoir être consommés par `_pdf_page_text` et
                # `_extraire_lignes_periodes_depuis_page_pdf`.
                mots = []
                for texte_mot, gauche, largeur, haut, hauteur_mot in zip(
                    donnees["text"],
                    donnees["left"],
                    donnees["width"],
                    donnees["top"],
                    donnees["height"],
                ):
                    propre = (texte_mot or "").strip()
                    if not propre:
                        continue
                    gauche_f = float(gauche)
                    haut_f = float(haut)
                    mots.append({
                        "text": propre,
                        "x0": gauche_f,
                        "x1": gauche_f + float(largeur),
                        "top": haut_f,
                        "bottom": haut_f + float(hauteur_mot),
                    })
                resultats[i] = (texte, mots)
    except Exception:
        return {}

    return resultats


def _ocr_pdf_octets(octets_bruts: bytes, langues: str = "fra+eng"):
    """OCR complet d'un PDF via `_ocr_pages`.

    Wrapper qui concatène les textes de toutes les pages et retourne
    aussi les mots coordonnés par page, pour la stratégie d'extraction
    par coordonnées (`_extraire_lignes_periodes_depuis_page_pdf`).

    Retourne `("", [])` si Tesseract n'est pas disponible ou si l'OCR
    échoue.
    """
    pages = _ocr_pages(
        octets_bruts, indices_cibles=None, langues=langues, avec_mots=True
    )
    if not pages:
        return ("", [])

    textes_pages = []
    mots_pages = []
    for i in sorted(pages.keys()):
        texte, mots = pages[i]
        textes_pages.append(texte)
        mots_pages.append(mots)
    return ("\n".join(textes_pages), mots_pages)


class _PageOCR:
    """Imite une page pdfplumber pour `_extraire_lignes_periodes_depuis_page_pdf`.

    Ne gère que `extract_words(...)` ; les tableaux structurés n'existent
    pas dans un scan. Les clés `text`/`top`/`x0`/`x1` reproduisent celles
    de pdfplumber, car la méthode d'extraction par coordonnées est écrite
    autour de cette API.
    """

    def __init__(self, mots):
        self._mots = mots

    def extract_words(self, **_kwargs):
        return list(self._mots)


def _construire_dataframe(lignes):
    """Crée un DataFrame à partir des lignes extraites, en sécurité."""
    if not lignes:
        return pd.DataFrame()

    largeur = max(len(ligne) for ligne in lignes)
    lignes_normalisees = [
        list(ligne) + [None] * (largeur - len(ligne))
        for ligne in lignes
    ]

    entetes = []
    for index, valeur in enumerate(lignes_normalisees[0]):
        nom = str(valeur).strip() if valeur is not None else ""
        entetes.append(nom or f"Colonne_{index + 1}")

    vues = {}
    entetes_uniques = []
    for nom in entetes:
        vues[nom] = vues.get(nom, 0) + 1
        entetes_uniques.append(
            nom if vues[nom] == 1 else f"{nom}_{vues[nom]}"
        )

    return pd.DataFrame(lignes_normalisees[1:], columns=entetes_uniques)


def detecter_separateur_decimal(texte: str) -> str:
    """Détecte le séparateur décimal utilisé par un document.

    Retourne ',' pour un document francophone (SCE tunisien, PCG français,
    ex. « 1 281 772 854,965 »), ou '.' pour un document anglophone
    (IFRS, ex. « 1,281,772,854.965 »).

    Détection basée sur des signaux clairs :
    - Au moins un nombre avec plusieurs virgules séparées par des groupes
      de 3 chiffres (« 12,345,678 ») → anglais.
    - Au moins un nombre virgule(s)-puis-point décimal
      (« 12,345.67 ») → anglais.
    - Au moins un nombre espace(s)-de-milliers puis virgule décimale
      (« 1 281,965 ») → français.

    Défaut : français (SCE tunisien).
    """
    if not texte:
        return ","
    # Signal fort pour l'anglais : plusieurs virgules séparant des groupes
    # de 3 chiffres.
    if re.search(r"\d{1,3}(?:,\d{3}){2,}\b", texte):
        return "."
    # Signal fort pour l'anglais : virgule(s) puis point décimal.
    if re.search(r"\d{1,3}(?:,\d{3})+\.\d+", texte):
        return "."
    # Signal fort pour le français : espace de milliers + virgule décimale.
    if re.search(r"\d(?:[\s  ]\d{3})+,\d+", texte):
        return ","
    return ","


def parser_montant(valeur, separateur_decimal: str = ","):
    """Analyse les formats numériques financiers français et anglais.

    `separateur_decimal` indique le séparateur décimal attendu pour les
    cas ambigus (ex. « 1,250 » = 1250 en anglais vs 1,250 en français).
    Valeurs acceptées : ',' (français — défaut) ou '.' (anglais).
    """
    if valeur is None:
        return None

    s = str(valeur).strip()

    if not s or s.lower() in {
        "nan", "none", "n/a", "na", "-", "—", "–"
    }:
        return None

    negatif = s.startswith("(") and s.endswith(")")

    s = (
        s.replace("\u202f", "")
        .replace("\xa0", "")
        .replace(" ", "")
    )

    s = re.sub(
        r"(?i)(TND|DT|EUR|USD|MAD|GBP|CHF|€|\$)",
        "",
        s,
    )

    s = s.replace("%", "").replace("(", "").replace(")", "")

    signe = -1 if s.startswith("-") else 1
    if s.startswith(("-", "+")):
        s = s[1:]

    if not s or not re.fullmatch(r"\d[\d.,]*", s):
        return None

    if "," in s and "." in s:
        # Les deux séparateurs présents : le dernier rencontré est la
        # décimale (règle robuste, indépendante de la locale).
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")

    elif "," in s:
        parties = s.split(",")
        if len(parties) == 2:
            droite = parties[1]
            if separateur_decimal == ",":
                # Document francophone : la virgule est décimale.
                # Ex. « 854,965 » = 854.965 ; « 3,5 » = 3.5.
                s = parties[0] + "." + droite
            else:
                # Document anglophone : exactement 3 chiffres après la
                # virgule → séparateur de milliers (« 1,250 » = 1250).
                # Sinon décimale « légère » (« 1,25 » = 1.25).
                if len(droite) == 3:
                    s = parties[0] + droite
                else:
                    s = parties[0] + "." + droite
        else:
            # Plusieurs virgules : séparateur de milliers
            # (ex. « 1,234,567 »).
            s = "".join(parties)

    elif s.count(".") > 1:
        parties = s.split(".")
        s = "".join(parties[:-1]) + "." + parties[-1]

    elif "." in s:
        gauche, droite = s.split(".", 1)
        # Un point suivi d'exactement 3 chiffres : séparateur de milliers
        # en français (« 1.250 » = 1250), décimal en anglais.
        if (
            separateur_decimal == ","
            and len(droite) == 3
            and gauche.isdigit()
            and droite.isdigit()
        ):
            s = gauche + droite

    try:
        resultat = float(s)
        return -resultat if negatif else signe * resultat
    except ValueError:
        return None


def _parser_date(valeur):
    """Convertit une chaîne de date supportée en objet date."""
    correspondance = MOTIF_DATE.search(str(valeur))
    if not correspondance:
        return None

    try:
        return datetime.strptime(correspondance.group(), "%d/%m/%Y").date()
    except ValueError:
        return None


# Regex d'un montant financier à la française :
# - entier avec groupes de 3 chiffres séparés par espace / espace fine /
#   insécable, optionnellement signé, optionnellement entre parenthèses,
# - partie décimale avec virgule ou point.
_MOTIF_MONTANT_FR = re.compile(
    r"-?\(?-?\d{1,3}(?:[\s \xa0]\d{3})+(?:[.,]\d+)?\)?"
    r"|-?\(?-?\d+(?:[.,]\d+)?\)?"
)


def _decouper_montants_dans_ligne(ligne, nb_periodes):
    """Essaie d'identifier les nb_periodes derniers montants d'une ligne
    dont les colonnes sont séparées par un simple espace.

    Retourne la liste des chaînes montants (ordre d'apparition) ou None
    si la ligne ne se prête pas à cette interprétation.
    """
    if nb_periodes < 2:
        return None
    correspondances = list(_MOTIF_MONTANT_FR.finditer(ligne))
    if len(correspondances) < nb_periodes:
        return None
    # Si la fin de la ligne n'est pas un montant (ex. phrase explicative),
    # ce découpage est risqué : on l'évite.
    queue = ligne[correspondances[-1].end():].strip()
    if queue and re.search(r"[A-Za-zÀ-ÿ]", queue):
        return None
    # Prendre les N derniers montants : ce sont les colonnes de période.
    correspondances_periode = correspondances[-nb_periodes:]
    # Vérifier que les montants « colonnes » sont consécutifs : entre
    # chaque, il n'y a que des espaces (pas de texte intermédiaire).
    for a, b in zip(correspondances_periode, correspondances_periode[1:]):
        entre = ligne[a.end():b.start()]
        if entre.strip():
            return None
    return [m.group(0) for m in correspondances_periode]


def _extraire_lignes_periodes_depuis_texte(texte, separateur_decimal: str = ","):
    """
    Extrait les lignes contenant plusieurs périodes financières.

    Exemple :
    Total des actifs courants    2 400    2 500

    quand l'en-tête précédent contient 31/12/2025 et 31/12/2024.
    """
    enregistrements = []
    periodes_courantes = []

    for ligne_originale in texte.splitlines():
        ligne = ligne_originale.strip()

        if not ligne:
            continue

        dates = MOTIF_DATE.findall(ligne)

        # Un en-tête contenant plusieurs dates définit les colonnes.
        if len(dates) >= 2:
            periodes_courantes = dates
            continue

        if not periodes_courantes:
            continue

        # Le texte PDF sépare souvent les colonnes par plusieurs espaces.
        # Un espace unique à l'intérieur de « 3 600 » est préservé comme séparateur de milliers.
        parties = [
            partie.strip()
            for partie in re.split(r"\s{2,}", ligne)
            if partie.strip()
        ]

        libelle = None
        cellules_montants = None

        if len(parties) >= len(periodes_courantes) + 1:
            # Chemin standard : au moins N+1 parties « bien séparées ».
            libelle = parties[0].strip(" •-\t")
            cellules_montants = parties[-len(periodes_courantes):]
        else:
            # Chemin de repli : colonnes collées par un simple espace
            # (fréquent avec pdfplumber sur certains PDF réels).
            queue = _decouper_montants_dans_ligne(ligne, len(periodes_courantes))
            if queue is None:
                continue
            # Le libellé est le texte avant le premier montant de queue.
            premiere_correspondance_queue = next(
                m for m in _MOTIF_MONTANT_FR.finditer(ligne)
                if m.group(0) == queue[0]
                and m.start() >= ligne.rfind(queue[-1]) - 500
            )
            texte_libelle = ligne[:premiere_correspondance_queue.start()].strip(" •-\t ")
            # Retirer un éventuel numéro de note isolé en fin de libellé
            # (ex. « Revenus 29 ... » → « Revenus »).
            texte_libelle = re.sub(
                r"[\s ]+-?\(?\d{1,3}\)?[\s ]*$",
                "",
                texte_libelle,
            ).strip()
            if not texte_libelle:
                continue
            libelle = texte_libelle
            cellules_montants = queue

        if not libelle:
            continue
        if not re.search(r"[A-Za-zÀ-ÿ]", libelle):
            continue

        for periode, cellule in zip(periodes_courantes, cellules_montants):
            valeur = parser_montant(cellule, separateur_decimal=separateur_decimal)

            if valeur is None:
                continue

            enregistrements.append({
                "libelle": libelle,
                "valeur": valeur,
                "periode": periode,
                "ligne_source": ligne,
            })

    return enregistrements


def _extraire_lignes_periodes_depuis_tableaux(tableaux, separateur_decimal: str = ","):
    """Extrait les enregistrements depuis des DataFrames à colonnes datées."""
    enregistrements = []

    for df in tableaux:
        if df is None or df.empty:
            continue

        colonnes_periodes = []

        for colonne in df.columns:
            correspondance_periode = MOTIF_DATE.search(str(colonne))
            if correspondance_periode:
                colonnes_periodes.append(
                    (colonne, correspondance_periode.group())
                )

        if len(colonnes_periodes) < 2:
            continue

        for _, ligne in df.iterrows():
            # La première colonne contient en général le libellé comptable.
            libelle = str(ligne.iloc[0]).strip() if len(ligne) else ""

            if not libelle or libelle.lower() in {"nan", "none"}:
                continue

            if not re.search(r"[A-Za-zÀ-ÿ]", libelle):
                continue

            for colonne, periode in colonnes_periodes:
                valeur = parser_montant(ligne[colonne], separateur_decimal=separateur_decimal)

                if valeur is None:
                    continue

                enregistrements.append({
                    "libelle": libelle,
                    "valeur": valeur,
                    "periode": periode,
                    "ligne_source": f"{libelle} | {periode}",
                })

    return enregistrements


def _extraire_lignes_periodes_depuis_page_pdf(
    page, tolerance_y=3, separateur_decimal: str = ","
):
    """Extrait les lignes financières à partir des mots coordonnés et des
    colonnes datées comme points d'ancrage.

    `tolerance_y` : tolérance verticale pour regrouper les mots d'une
    même ligne. Les PDF natifs ont des glyphes alignés précisément (2-3
    pixels) ; une page OCR scannée à 300 dpi a besoin de bien plus
    (textbaselines légèrement décalés, 15-30 pixels).

    Partage le helper `_grouper_mots_par_ligne` avec `_pdf_page_text` :
    une seule logique de bucketing, deux usages distincts en aval.
    """
    lignes_mots = _grouper_mots_par_ligne(page, tolerance_y=tolerance_y)
    if not lignes_mots:
        return []

    enregistrements = []
    periodes = []
    ancres = []

    motif_date = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
    motif_nombre = re.compile(
        r"^\(?[-+]?\d[\d.,]*\)?%?$"
    )

    for mots_ligne in lignes_mots:

        # Détecter les dates et leur position horizontale.
        mots_dates = [
            mot for mot in mots_ligne
            if motif_date.fullmatch(mot["text"])
            and _parser_date(mot["text"]) is not None
        ]

        if len(mots_dates) >= 2:
            mots_dates.sort(key=lambda w: w["x0"])
            periodes = [mot["text"] for mot in mots_dates]
            ancres = [
                (mot["x0"] + mot["x1"]) / 2
                for mot in mots_dates
            ]
            continue

        if len(periodes) < 2 or len(ancres) < 2:
            continue

        mots_numeriques = [
            mot for mot in mots_ligne
            if motif_nombre.fullmatch(mot["text"])
        ]
        if not mots_numeriques:
            continue

        # Construire les groupes numériques. Deux stratégies selon la
        # densité des mots :
        #  (a) montants expansés (PDF natif qui livre « 148 », « 273 »,
        #      « 413 » séparément : nombreux mots, petits gaps à
        #      l'intérieur d'un montant, grands gaps entre colonnes) ;
        #  (b) montants compacts (OCR qui concatène « 148273413 » en un
        #      seul token : un mot = une colonne).
        # On distingue avec le nombre de mots par rapport au nombre de
        # colonnes attendues.
        gaps = [
            mots_numeriques[i + 1]["x0"] - mots_numeriques[i]["x1"]
            for i in range(len(mots_numeriques) - 1)
        ]

        if len(mots_numeriques) > len(periodes) and gaps:
            # Stratégie (a) : seuil adaptatif gap > 2.5 × médiane,
            # avec minimum à 6 pixels pour les PDF natifs très serrés.
            gaps_tries = sorted(gaps)
            gap_median = gaps_tries[len(gaps_tries) // 2]
            seuil_gap = max(6.0, gap_median * 2.5)

            groupes_numeriques = []
            groupe_courant = []
            for mot in mots_numeriques:
                if groupe_courant:
                    gap = mot["x0"] - groupe_courant[-1]["x1"]
                    if gap > seuil_gap:
                        groupes_numeriques.append(groupe_courant)
                        groupe_courant = []
                groupe_courant.append(mot)
            if groupe_courant:
                groupes_numeriques.append(groupe_courant)
        else:
            # Stratégie (b) : chaque mot numérique est un groupe
            # autonome (OCR qui regroupe déjà les chiffres par montant).
            groupes_numeriques = [[mot] for mot in mots_numeriques]

        # Il faut une valeur identifiable pour chaque période.
        if len(groupes_numeriques) < len(periodes):
            continue

        # Associer chaque groupe numérique à la date la plus proche en
        # partant de la DROITE : la colonne la plus à droite correspond
        # toujours à la période la plus à droite. Cela empêche qu'une
        # référence de note (petit nombre dans une colonne « Notes »
        # placée à GAUCHE des montants) soit assignée à une période par
        # proximité, ce qui écraserait la vraie valeur.
        assignes = {}
        ids_groupes_assignes = set()
        for index_ancre in sorted(range(len(ancres)), reverse=True):
            meilleur_groupe = None
            meilleure_distance = None
            for idx, groupe in enumerate(groupes_numeriques):
                if idx in ids_groupes_assignes:
                    continue
                centre = (groupe[0]["x0"] + groupe[-1]["x1"]) / 2
                distance = abs(centre - ancres[index_ancre])
                if meilleure_distance is None or distance < meilleure_distance:
                    meilleure_distance = distance
                    meilleur_groupe = (idx, groupe)
            if meilleur_groupe is not None:
                idx, groupe = meilleur_groupe
                assignes[index_ancre] = groupe
                ids_groupes_assignes.add(idx)

        if len(assignes) < len(periodes):
            continue

        premier_x_valeur = min(
            groupe[0]["x0"] for groupe in assignes.values()
        )

        mots_libelle = [
            mot["text"] for mot in mots_ligne
            if mot["x1"] <= premier_x_valeur
        ]

        libelle = " ".join(mots_libelle).strip(" :-•\t")
        libelle = re.sub(
            r"\([A-Za-z]\.?\d+(?:\.\d+)*\)",
            "",
            libelle,
        ).strip()
        # Retirer un numéro de note isolé en fin de libellé
        # (ex. « Revenus 29 » → « Revenus », « Stocks 10 » → « Stocks »).
        libelle = re.sub(
            r"[\s ]+-?\(?\d{1,3}\)?[\s ]*$",
            "",
            libelle,
        ).strip()

        if not libelle or not re.search(r"[A-Za-zÀ-ÿ]", libelle):
            continue

        for index, periode in enumerate(periodes):
            groupe = assignes.get(index)
            if not groupe:
                continue

            texte_montant = "".join(mot["text"] for mot in groupe)
            valeur = parser_montant(texte_montant, separateur_decimal=separateur_decimal)

            if valeur is None:
                continue

            enregistrements.append({
                "libelle": libelle,
                "valeur": valeur,
                "periode": periode,
                "ligne_source": " ".join(
                    mot["text"] for mot in mots_ligne
                ),
            })

    return enregistrements



def _dedoublonner_enregistrements(enregistrements):
    """Retire les doublons sans fusionner des périodes différentes."""
    resultat = []
    vus = set()

    for enreg in enregistrements:
        signature = (
            re.sub(r"\s+", " ", enreg["libelle"]).strip().casefold(),
            enreg["valeur"],
            enreg.get("periode"),
        )

        if signature in vus:
            continue

        vus.add(signature)
        resultat.append(enreg)

    return resultat


def extraire_document(fichier_charge):
    """Extrait le texte, les tableaux et les enregistrements avec périodes."""
    enregistrements_pdf = []
    suffixe = Path(fichier_charge.name).suffix.lower()

    if suffixe not in FORMATS_PRIS_EN_CHARGE:
        raise ValueError(f"Format non pris en charge : {suffixe}")

    octets_bruts = fichier_charge.getvalue()
    tampon = io.BytesIO(octets_bruts)

    texte = ""
    tableaux = []
    ocr_utilise = False

    if suffixe == ".txt":
        texte = octets_bruts.decode("utf-8-sig", errors="replace")

    elif suffixe == ".csv":
        try:
            df = pd.read_csv(
                tampon,
                sep=None,
                engine="python",
                dtype=str,
                keep_default_na=False,
            )
        except Exception:
            tampon.seek(0)
            df = pd.read_csv(
                tampon,
                dtype=str,
                keep_default_na=False,
            )

        if not df.empty:
            tableaux.append(df)
            texte = df.to_string(index=False)

    elif suffixe in {".xlsx", ".xls"}:
        feuilles = pd.read_excel(
            tampon,
            sheet_name=None,
            dtype=str,
        )

        morceaux_texte = []

        for nom_feuille, df in feuilles.items():
            if not df.empty:
                tableaux.append(df)
                morceaux_texte.append(
                    f"\nFeuille : {nom_feuille}\n{df.to_string(index=False)}"
                )

        texte = "\n".join(morceaux_texte)

    elif suffixe == ".pdf":
        # Deux passes : (1) assembler le texte et les tableaux pour
        # pouvoir détecter la locale, (2) rejouer l'extraction par
        # coordonnées avec le bon séparateur décimal.
        try:
            import pdfplumber

            pages_pdf = []  # pages pdfplumber
            with pdfplumber.open(io.BytesIO(octets_bruts)) as pdf:
                textes_pages = []

                for numero_page, page in enumerate(pdf.pages, start=1):
                    texte_page = page.extract_text() or ""
                    textes_pages.append(
                        f"\n--- Page {numero_page} ---\n{texte_page}"
                    )

                    # Extraction des tableaux structurés.
                    for tableau_extrait in page.extract_tables() or []:
                        df = _construire_dataframe(tableau_extrait)
                        if not df.empty:
                            tableaux.append(df)

                    pages_pdf.append(page)

                texte = "\n".join(textes_pages)
                separateur_decimal = detecter_separateur_decimal(texte)

                # Deuxième passe : parseur par coordonnées, avec la
                # locale correctement détectée.
                for page in pages_pdf:
                    enregistrements_pdf.extend(
                        _extraire_lignes_periodes_depuis_page_pdf(
                            page, separateur_decimal=separateur_decimal
                        )
                    )

        except Exception as exc:
            raise ValueError(
                f"Impossible de lire ce PDF : {exc}"
            ) from exc

        # OCR en repli si pdfplumber n'a rien renvoyé d'exploitable.
        # Déclencheur : trop peu de caractères par page ET aucun record
        # extrait des tableaux structurés ou des coordonnées.
        try:
            import pdfplumber  # noqa: F401
            with pdfplumber.open(io.BytesIO(octets_bruts)) as _pdf:
                nombre_pages = max(1, len(_pdf.pages))
        except Exception:
            nombre_pages = 1

        densite_texte = len(texte.strip()) / nombre_pages

        if (
            densite_texte < OCR_MIN_CARACTERES_PAR_PAGE
            and not enregistrements_pdf
            and not tableaux
        ):
            if ocr_disponible():
                texte_ocr, mots_pages_ocr = _ocr_pdf_octets(octets_bruts)
                if texte_ocr:
                    texte = nettoyer_texte_ocr(texte_ocr)
                    ocr_utilise = True
                    # La locale OCR dépend du contenu lu : on redétecte.
                    separateur_decimal = detecter_separateur_decimal(texte)
                    # Rejouer l'extraction par coordonnées sur les mots
                    # retournés par l'OCR : on retrouve ainsi les colonnes
                    # grâce aux positions x, y des mots, exactement comme
                    # pour un PDF à couche texte. Pages scannées à 300 dpi
                    # → tolerance_y plus permissif (lignes légèrement
                    # obliques, textbaselines imprécises).
                    for mots in mots_pages_ocr:
                        enregistrements_pdf.extend(
                            _extraire_lignes_periodes_depuis_page_pdf(
                                _PageOCR(mots),
                                tolerance_y=30,
                                separateur_decimal=separateur_decimal,
                            )
                        )

    elif suffixe == ".docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(octets_bruts))

            morceaux_texte = [
                paragraphe.text
                for paragraphe in doc.paragraphs
                if paragraphe.text.strip()
            ]

            for tableau_extrait in doc.tables:
                lignes = [
                    [cellule.text.strip() for cellule in ligne.cells]
                    for ligne in tableau_extrait.rows
                ]

                df = _construire_dataframe(lignes)
                if not df.empty:
                    tableaux.append(df)

            texte = "\n".join(morceaux_texte)

        except Exception as exc:
            raise ValueError(
                f"Impossible de lire ce document Word : {exc}"
            ) from exc

    # Détection de la locale finale sur tout le texte extrait.
    # Pour un PDF, elle a déjà été utilisée dans les deux passes
    # ci-dessus ; on la calcule à nouveau ici pour couvrir TXT/CSV/
    # XLSX/DOCX et pour l'exposer dans le résultat.
    separateur_decimal = detecter_separateur_decimal(texte)

    # Préférer les enregistrements de tableaux structurés, puis la mise
    # en page extraite du PDF/texte. Les doublons libellé/valeur/période
    # sont retirés.

    enregistrements_tableaux = _extraire_lignes_periodes_depuis_tableaux(
        tableaux, separateur_decimal=separateur_decimal
    )

    if suffixe == ".pdf":
        # Priorité aux tableaux structurés et aux coordonnées.
        # Le texte sert de repli si aucune extraction structurée n'a réussi.
        if enregistrements_tableaux or enregistrements_pdf:
            enregistrements = _dedoublonner_enregistrements(
                enregistrements_tableaux + enregistrements_pdf
            )
        else:
            enregistrements = _dedoublonner_enregistrements(
                _extraire_lignes_periodes_depuis_texte(
                    texte, separateur_decimal=separateur_decimal
                )
            )
    else:
        enregistrements_texte = _extraire_lignes_periodes_depuis_texte(
            texte, separateur_decimal=separateur_decimal
        )
        enregistrements = _dedoublonner_enregistrements(
            enregistrements_tableaux + enregistrements_texte
        )


    if ocr_utilise:
        for enreg in enregistrements:
            enreg["ocr"] = True

    return {
        "texte": texte,
        "tableaux": tableaux,
        "enregistrements": enregistrements,
        "suffixe": suffixe,
        "ocr_utilise": ocr_utilise,
        "ocr_disponible": ocr_disponible(),
        "separateur_decimal": separateur_decimal,
    }


def extraire_valeurs_etiquetees(texte, separateur_decimal: str = None):
    """Extrait des couples libellé/valeur en préservant les périodes détectées.

    `separateur_decimal` : si None, détection automatique à partir de
    tout le texte. Sinon ',' (français) ou '.' (anglais) force la
    locale.
    """
    if separateur_decimal is None:
        separateur_decimal = detecter_separateur_decimal(texte)

    enregistrements_periodes = _extraire_lignes_periodes_depuis_texte(
        texte, separateur_decimal=separateur_decimal
    )

    if enregistrements_periodes:
        return _dedoublonner_enregistrements(enregistrements_periodes)

    resultats = []
    vus = set()

    devise = r"(?:TND|DT|EUR|USD|MAD|GBP|CHF|€|\$)"

    motifs = [
        re.compile(r"^(.*?)\s*:\s*(.*?)\s*$"),
        re.compile(
            r"^(.*?)\s+"
            r"([-+]?\(?\d[\d\s.,]*\)?)"
            rf"\s*{devise}?\s*$",
            re.IGNORECASE,
        ),
    ]

    for ligne_originale in texte.splitlines():
        ligne = ligne_originale.strip()

        if not ligne or len(ligne) > 500:
            continue

        libelle = ""
        montant = None

        correspondance_colon = motifs[0].match(ligne)

        if correspondance_colon:
            libelle_possible = correspondance_colon.group(1).strip(" •-\t")
            valeur_possible = correspondance_colon.group(2).strip()
            parsed = parser_montant(
                valeur_possible, separateur_decimal=separateur_decimal
            )

            if libelle_possible and parsed is not None:
                libelle = libelle_possible
                montant = parsed

        if montant is None:
            correspondance = motifs[1].match(ligne)

            if correspondance:
                libelle_possible = correspondance.group(1).strip(" •-\t")
                valeur_possible = correspondance.group(2).strip()
                parsed = parser_montant(
                    valeur_possible, separateur_decimal=separateur_decimal
                )

                if (
                    libelle_possible
                    and re.search(r"[A-Za-zÀ-ÿ]", libelle_possible)
                    and parsed is not None
                ):
                    libelle = libelle_possible
                    montant = parsed

        if montant is None or not libelle:
            continue

        signature = (libelle.casefold(), montant)

        if signature not in vus:
            vus.add(signature)
            resultats.append({
                "libelle": libelle,
                "valeur": montant,
                "periode": None,
                "ligne_source": ligne,
            })

    return resultats
