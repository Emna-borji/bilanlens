from pathlib import Path
from datetime import datetime
import io
import re

import pandas as pd

SUPPORTED = {".pdf", ".xlsx", ".xls", ".csv", ".txt", ".docx"}

DATE_PATTERN = re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b")


def _make_dataframe(rows):
    """Create a DataFrame safely from extracted rows."""
    if not rows:
        return pd.DataFrame()

    width = max(len(row) for row in rows)
    normalized_rows = [
        list(row) + [None] * (width - len(row))
        for row in rows
    ]

    headers = []
    for index, value in enumerate(normalized_rows[0]):
        name = str(value).strip() if value is not None else ""
        headers.append(name or f"Colonne_{index + 1}")

    seen = {}
    unique_headers = []
    for name in headers:
        seen[name] = seen.get(name, 0) + 1
        unique_headers.append(
            name if seen[name] == 1 else f"{name}_{seen[name]}"
        )

    return pd.DataFrame(normalized_rows[1:], columns=unique_headers)


def parse_amount(value):
    """Parse common French and English financial number formats."""
    if value is None:
        return None

    s = str(value).strip()

    if not s or s.lower() in {
        "nan", "none", "n/a", "na", "-", "—", "–"
    }:
        return None

    negative = s.startswith("(") and s.endswith(")")

    s = (
        s.replace("\u202f", "")
        .replace("\xa0", "")
        .replace(" ", "")
    )

    s = re.sub(
        r"(?i)(TND|DT|EUR|USD|MAD|GBP|CHF|€|\$)",
        "",
        s,
    )

    s = s.replace("%", "").replace("(", "").replace(")", "")

    sign = -1 if s.startswith("-") else 1
    if s.startswith(("-", "+")):
        s = s[1:]

    if not s or not re.fullmatch(r"\d[\d.,]*", s):
        return None

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")

    elif "," in s:
        parts = s.split(",")
        if len(parts) == 2 and len(parts[-1]) in (1, 2):
            s = parts[0] + "." + parts[-1]
        else:
            s = "".join(parts)

    elif s.count(".") > 1:
        parts = s.split(".")
        s = "".join(parts[:-1]) + "." + parts[-1]

    elif "." in s:
        left, right = s.split(".", 1)
        if len(right) == 3 and left.isdigit() and right.isdigit():
            s = left + right

    try:
        result = float(s)
        return -result if negative else sign * result
    except ValueError:
        return None


def _parse_date(value):
    """Convert a supported date string to a date object."""
    match = DATE_PATTERN.search(str(value))
    if not match:
        return None

    try:
        return datetime.strptime(match.group(), "%d/%m/%Y").date()
    except ValueError:
        return None


def _extract_period_records_from_text(text):
    """
    Extract rows containing multiple financial periods.

    Example:
    Total des actifs courants    2 400    2 500

    when the preceding header contains 31/12/2025 and 31/12/2024.
    """
    records = []
    current_periods = []

    for original_line in text.splitlines():
        line = original_line.strip()

        if not line:
            continue

        dates = DATE_PATTERN.findall(line)

        # A header containing multiple dates defines the column periods.
        if len(dates) >= 2:
            current_periods = dates
            continue

        if not current_periods:
            continue

        # PDF text often separates table columns with multiple spaces.
        # A single space inside "3 600" is preserved as a thousands separator.
        parts = [
            part.strip()
            for part in re.split(r"\s{2,}", line)
            if part.strip()
        ]

        if len(parts) < len(current_periods) + 1:
            continue

        label = parts[0].strip(" •-\t")
        if not re.search(r"[A-Za-zÀ-ÿ]", label):
            continue

        # The final cells correspond to the date columns.
        amount_cells = parts[-len(current_periods):]

        for period, amount_cell in zip(current_periods, amount_cells):
            value = parse_amount(amount_cell)

            if value is None:
                continue

            records.append({
                "label": label,
                "value": value,
                "period": period,
                "source_line": line,
            })

    return records


def _extract_period_records_from_tables(tables):
    """Extract records directly from DataFrames with date columns."""
    records = []

    for df in tables:
        if df is None or df.empty:
            continue

        period_columns = []

        for column in df.columns:
            period_match = DATE_PATTERN.search(str(column))
            if period_match:
                period_columns.append(
                    (column, period_match.group())
                )

        if len(period_columns) < 2:
            continue

        for _, row in df.iterrows():
            # Usually the first column contains the accounting label.
            label = str(row.iloc[0]).strip() if len(row) else ""

            if not label or label.lower() in {"nan", "none"}:
                continue

            if not re.search(r"[A-Za-zÀ-ÿ]", label):
                continue

            for column, period in period_columns:
                value = parse_amount(row[column])

                if value is None:
                    continue

                records.append({
                    "label": label,
                    "value": value,
                    "period": period,
                    "source_line": f"{label} | {period}",
                })

    return records


def _extract_period_records_from_pdf_page(page):
    """Extract financial rows using word coordinates and date-column anchors."""
    words = page.extract_words(
        x_tolerance=2,
        y_tolerance=3,
        keep_blank_chars=False,
    )

    if not words:
        return []

    # Regrouper les mots par ligne visuelle.
    lines = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        target = None

        for line in reversed(lines[-3:]):
            if abs(line["top"] - word["top"]) <= 3:
                target = line
                break

        if target is None:
            lines.append({"top": word["top"], "words": [word]})
        else:
            target["words"].append(word)

    lines.sort(key=lambda line: line["top"])

    records = []
    periods = []
    anchors = []

    date_re = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
    number_re = re.compile(
        r"^\(?[-+]?\d[\d.,]*\)?%?$"
    )

    for line in lines:
        row_words = sorted(line["words"], key=lambda w: w["x0"])

        # Détecter les dates et leur position horizontale.
        date_words = [
            word for word in row_words
            if date_re.fullmatch(word["text"])
            and _parse_date(word["text"]) is not None
        ]

        if len(date_words) >= 2:
            date_words.sort(key=lambda w: w["x0"])
            periods = [word["text"] for word in date_words]
            anchors = [
                (word["x0"] + word["x1"]) / 2
                for word in date_words
            ]
            continue

        if len(periods) < 2 or len(anchors) < 2:
            continue

        # Construire les groupes numériques à partir des coordonnées.
        numeric_groups = []
        current_group = []

        for word in row_words:
            token = word["text"]

            if not number_re.fullmatch(token):
                continue

            if current_group:
                gap = word["x0"] - current_group[-1]["x1"]

                # Les morceaux d'un montant comme "3 600"
                # sont proches ; les colonnes sont normalement plus éloignées.
                if gap > 6:
                    numeric_groups.append(current_group)
                    current_group = []

            current_group.append(word)

        if current_group:
            numeric_groups.append(current_group)

        # Il faut une valeur identifiable pour chaque période.
        if len(numeric_groups) < len(periods):
            continue

        # Associer chaque groupe numérique à la date la plus proche.
        assigned = {}
        for group in numeric_groups:
            center = (group[0]["x0"] + group[-1]["x1"]) / 2
            period_index = min(
                range(len(anchors)),
                key=lambda i: abs(center - anchors[i]),
            )

            if period_index not in assigned:
                assigned[period_index] = group

        if len(assigned) < len(periods):
            continue

        first_value_x = min(
            group[0]["x0"] for group in assigned.values()
        )

        label_words = [
            word["text"] for word in row_words
            if word["x1"] <= first_value_x
        ]

        label = " ".join(label_words).strip(" :-•\t")
        label = re.sub(
            r"\([A-Za-z]\.?\d+(?:\.\d+)*\)",
            "",
            label,
        ).strip()

        if not label or not re.search(r"[A-Za-zÀ-ÿ]", label):
            continue

        for index, period in enumerate(periods):
            group = assigned.get(index)
            if not group:
                continue

            amount_text = "".join(word["text"] for word in group)
            value = parse_amount(amount_text)

            if value is None:
                continue

            records.append({
                "label": label,
                "value": value,
                "period": period,
                "source_line": " ".join(
                    word["text"] for word in row_words
                ),
            })

    return records



def _deduplicate_records(records):
    """Remove duplicates without merging different periods."""
    result = []
    seen = set()

    for record in records:
        signature = (
            re.sub(r"\s+", " ", record["label"]).strip().casefold(),
            record["value"],
            record.get("period"),
        )

        if signature in seen:
            continue

        seen.add(signature)
        result.append(record)

    return result


def extract_document(uploaded_file):
    """Extract text, tables and period-aware financial records."""
    pdf_records = []
    suffix = Path(uploaded_file.name).suffix.lower()

    if suffix not in SUPPORTED:
        raise ValueError(f"Format non pris en charge : {suffix}")

    raw = uploaded_file.getvalue()
    buffer = io.BytesIO(raw)

    text = ""
    tables = []

    if suffix == ".txt":
        text = raw.decode("utf-8-sig", errors="replace")

    elif suffix == ".csv":
        try:
            df = pd.read_csv(
                buffer,
                sep=None,
                engine="python",
                dtype=str,
                keep_default_na=False,
            )
        except Exception:
            buffer.seek(0)
            df = pd.read_csv(
                buffer,
                dtype=str,
                keep_default_na=False,
            )

        if not df.empty:
            tables.append(df)
            text = df.to_string(index=False)

    elif suffix in {".xlsx", ".xls"}:
        sheets = pd.read_excel(
            buffer,
            sheet_name=None,
            dtype=str,
        )

        text_parts = []

        for sheet_name, df in sheets.items():
            if not df.empty:
                tables.append(df)
                text_parts.append(
                    f"\nFeuille : {sheet_name}\n{df.to_string(index=False)}"
                )

        text = "\n".join(text_parts)

    
    elif suffix == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                page_texts = []

                for page_number, page in enumerate(pdf.pages, start=1):
                    page_texts.append(
                        f"\n--- Page {page_number} ---\n"
                        f"{page.extract_text() or ''}"
                    )

                    # Extraction des tableaux structurés.
                    for extracted_table in page.extract_tables() or []:
                        df = _make_dataframe(extracted_table)
                        if not df.empty:
                            tables.append(df)

                    # Extraction indépendante grâce aux coordonnées.
                    pdf_records.extend(
                        _extract_period_records_from_pdf_page(page)
                    )

                text = "\n".join(page_texts)

        except Exception as exc:
            raise ValueError(
                f"Impossible de lire ce PDF : {exc}"
            ) from exc


    elif suffix == ".docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(raw))

            text_parts = [
                paragraph.text
                for paragraph in doc.paragraphs
                if paragraph.text.strip()
            ]

            for extracted_table in doc.tables:
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in extracted_table.rows
                ]

                df = _make_dataframe(rows)
                if not df.empty:
                    tables.append(df)

            text = "\n".join(text_parts)

        except Exception as exc:
            raise ValueError(
                f"Impossible de lire ce document Word : {exc}"
            ) from exc

    # Prefer structured table records, then use the extracted PDF/text layout.
    # Duplicate label/value/period records are removed.

    table_records = _extract_period_records_from_tables(tables)

    if suffix == ".pdf":
        # Priorité aux tableaux structurés et aux coordonnées.
        # Le texte sert de repli si aucune extraction structurée n'a réussi.
        if table_records or pdf_records:
            records = _deduplicate_records(table_records + pdf_records)
        else:
            records = _deduplicate_records(
                _extract_period_records_from_text(text)
            )
    else:
        text_records = _extract_period_records_from_text(text)
        records = _deduplicate_records(table_records + text_records)


    return {
        "text": text,
        "tables": tables,
        "records": records,
        "suffix": suffix,
    }


def extract_labeled_values(text):
    """Extract label/value pairs, preserving periods when detected."""
    period_records = _extract_period_records_from_text(text)

    if period_records:
        return _deduplicate_records(period_records)

    results = []
    seen = set()

    currency = r"(?:TND|DT|EUR|USD|MAD|GBP|CHF|€|\$)"

    patterns = [
        re.compile(r"^(.*?)\s*:\s*(.*?)\s*$"),
        re.compile(
            r"^(.*?)\s+"
            r"([-+]?\(?\d[\d\s.,]*\)?)"
            rf"\s*{currency}?\s*$",
            re.IGNORECASE,
        ),
    ]

    for original_line in text.splitlines():
        line = original_line.strip()

        if not line or len(line) > 500:
            continue

        label = ""
        amount = None

        colon_match = patterns[0].match(line)

        if colon_match:
            possible_label = colon_match.group(1).strip(" •-\t")
            possible_value = colon_match.group(2).strip()
            parsed = parse_amount(possible_value)

            if possible_label and parsed is not None:
                label = possible_label
                amount = parsed

        if amount is None:
            match = patterns[1].match(line)

            if match:
                possible_label = match.group(1).strip(" •-\t")
                possible_value = match.group(2).strip()
                parsed = parse_amount(possible_value)

                if (
                    possible_label
                    and re.search(r"[A-Za-zÀ-ÿ]", possible_label)
                    and parsed is not None
                ):
                    label = possible_label
                    amount = parsed

        if amount is None or not label:
            continue

        signature = (label.casefold(), amount)

        if signature not in seen:
            seen.add(signature)
            results.append({
                "label": label,
                "value": amount,
                "period": None,
                "source_line": line,
            })

    return results